from __future__ import annotations

import os
import json
import queue
import subprocess
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk


BASE_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
BUILD_VERSION = "2026.09.20"
DEFAULT_XLSX = BASE_DIR / "data" / "customers.xlsx"
DEFAULT_CSV = BASE_DIR / "data" / "customers.csv"
GUI_STATE_FILE = BASE_DIR / "gui_state.json"


def default_input_path() -> Path:
    """Restore the last selected table when it still exists."""
    try:
        state = json.loads(GUI_STATE_FILE.read_text(encoding="utf-8"))
        saved = Path(str(state.get("last_input", "")).strip()).expanduser()
        if saved.is_file():
            return saved
    except (OSError, ValueError, TypeError):
        pass
    if DEFAULT_XLSX.exists():
        return DEFAULT_XLSX
    return DEFAULT_CSV


def save_input_path(path: Path) -> None:
    if not path.is_file():
        return
    try:
        GUI_STATE_FILE.write_text(
            json.dumps({"last_input": str(path.resolve())}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except OSError:
        # Remembering the path is convenience state; a read-only folder must
        # not prevent the robot from running.
        pass


def input_row_count(path: Path) -> int:
    """Preview the selected table without importing the browser worker."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            next(rows, None)  # header
            return sum(1 for row in rows if any(row))
        finally:
            workbook.close()
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return max(sum(1 for _ in handle) - 1, 0)


class RobotGui(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(f"瑞典 NPA 管理员自动化 v{BUILD_VERSION}")
        self.geometry("900x650")
        self.minsize(760, 520)
        self.process: subprocess.Popen[str] | None = None
        self.output_queue: queue.Queue[str] = queue.Queue()
        self.reader_thread: threading.Thread | None = None

        self.input_var = tk.StringVar(value=str(default_input_path()))
        self.headed_var = tk.BooleanVar(value=True)
        self.status_var = tk.StringVar(value="未运行")
        self._build_ui()
        self._poll_output()
        self._preview_input()

        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def _build_ui(self) -> None:
        self.columnconfigure(0, weight=1)
        self.rowconfigure(3, weight=1)

        title = ttk.Label(self, text="瑞典 NPA 管理员自动化", font=("Segoe UI", 18, "bold"))
        title.grid(row=0, column=0, padx=18, pady=(16, 8), sticky="w")

        input_frame = ttk.LabelFrame(self, text="客户表格")
        input_frame.grid(row=1, column=0, padx=18, pady=6, sticky="ew")
        input_frame.columnconfigure(1, weight=1)
        ttk.Label(input_frame, text="文件：").grid(row=0, column=0, padx=(10, 4), pady=10)
        input_entry = ttk.Entry(input_frame, textvariable=self.input_var)
        input_entry.grid(row=0, column=1, padx=4, pady=10, sticky="ew")
        ttk.Button(input_frame, text="选择文件", command=self._choose_input).grid(row=0, column=2, padx=4, pady=10)
        ttk.Button(input_frame, text="刷新预览", command=self._preview_input).grid(row=0, column=3, padx=(4, 10), pady=10)
        self.preview_label = ttk.Label(input_frame, text="")
        self.preview_label.grid(row=1, column=1, columnspan=3, padx=4, pady=(0, 8), sticky="w")

        options = ttk.LabelFrame(self, text="运行设置")
        options.grid(row=2, column=0, padx=18, pady=6, sticky="ew")
        ttk.Label(options, text="执行模式：正式执行（会点击 Add 和 Save）").grid(row=0, column=0, columnspan=3, padx=(10, 4), pady=10, sticky="w")
        ttk.Checkbutton(options, text="显示浏览器窗口", variable=self.headed_var).grid(row=1, column=1, padx=4, pady=(0, 10), sticky="w")
        ttk.Label(options, text="验证码：自动读取").grid(row=1, column=2, padx=4, pady=(0, 10), sticky="w")

        log_frame = ttk.LabelFrame(self, text="运行日志")
        log_frame.grid(row=3, column=0, padx=18, pady=6, sticky="nsew")
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)
        self.log_text = tk.Text(log_frame, wrap="word", state="disabled", font=("Consolas", 10))
        self.log_text.grid(row=0, column=0, padx=(8, 0), pady=8, sticky="nsew")
        scrollbar = ttk.Scrollbar(log_frame, orient="vertical", command=self.log_text.yview)
        scrollbar.grid(row=0, column=1, padx=(0, 8), pady=8, sticky="ns")
        self.log_text.configure(yscrollcommand=scrollbar.set)

        bottom = ttk.Frame(self)
        bottom.grid(row=4, column=0, padx=18, pady=(6, 16), sticky="ew")
        bottom.columnconfigure(0, weight=1)
        self.start_button = ttk.Button(bottom, text="开始运行", command=self._start)
        self.start_button.grid(row=0, column=1, padx=4)
        self.stop_button = ttk.Button(bottom, text="停止", command=self._stop, state="disabled")
        self.stop_button.grid(row=0, column=2, padx=4)
        ttk.Button(bottom, text="打开结果目录", command=self._open_runs).grid(row=0, column=3, padx=(4, 0))
        ttk.Label(bottom, textvariable=self.status_var).grid(row=1, column=0, columnspan=4, pady=(10, 0), sticky="w")

    def _append_log(self, text: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", text)
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _choose_input(self) -> None:
        chosen = filedialog.askopenfilename(
            title="选择客户表格",
            initialdir=str(BASE_DIR / "data"),
            filetypes=[("Excel 文件", "*.xlsx"), ("CSV 文件", "*.csv"), ("所有文件", "*.*")],
        )
        if chosen:
            self.input_var.set(chosen)
            save_input_path(Path(chosen))
            self._preview_input()

    def _preview_input(self) -> None:
        path = Path(self.input_var.get().strip())
        if not path.exists():
            self.preview_label.configure(text="文件不存在")
            return
        save_input_path(path)
        try:
            count = input_row_count(path)
            self.preview_label.configure(text=f"已读取 {count} 条记录；运行时会跳过 enabled=false/否/跳过 的记录。")
        except Exception as exc:
            self.preview_label.configure(text=f"读取失败：{exc}")

    def _start(self) -> None:
        if self.process is not None and self.process.poll() is None:
            return
        input_path = Path(self.input_var.get().strip())
        if not input_path.exists():
            messagebox.showerror("无法启动", "客户表格不存在，请先选择 .xlsx 或 .csv 文件。")
            return
        save_input_path(input_path)
        confirmed = messagebox.askyesno("确认正式执行", "任务会在 NPA 页面点击 Add，并在 Invoice email 页面点击 Save。确认继续吗？")
        if not confirmed:
            return
        if getattr(sys, "frozen", False):
            worker = BASE_DIR / "robot_worker.exe"
            if not worker.is_file():
                bundle_dir = Path(getattr(sys, "_MEIPASS", BASE_DIR))
                worker = bundle_dir / "runtime" / "robot_worker.exe"
            if not worker.is_file():
                messagebox.showerror("无法启动", "缺少后台执行组件，请重新获取完整的 RobotAdmin.exe。")
                return
            command = [str(worker), "--input", str(input_path)]
            config_path = BASE_DIR / "config.json"
            if not config_path.is_file():
                config_path = Path(getattr(sys, "_MEIPASS", BASE_DIR)) / "config.json"
            if config_path.is_file():
                command.extend(["--config", str(config_path)])
        else:
            command = [sys.executable, "-u", str(BASE_DIR / "main.py"), "--input", str(input_path)]
        if not self.headed_var.get():
            command.append("--headless")
        else:
            command.append("--headed")
        command.append("--auto-code")
        env = os.environ.copy()
        env["PYTHONUTF8"] = "1"
        env["PYTHONIOENCODING"] = "utf-8"
        try:
            popen_kwargs = {}
            if getattr(sys, "frozen", False):
                popen_kwargs["creationflags"] = getattr(subprocess, "CREATE_NO_WINDOW", 0)
            self.process = subprocess.Popen(
                command,
                cwd=str(BASE_DIR),
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding="utf-8",
                errors="replace",
                bufsize=1,
                env=env,
                **popen_kwargs,
            )
        except Exception as exc:
            messagebox.showerror("无法启动", str(exc))
            return
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")
        self._append_log("已启动任务。程序将自动读取并提交验证码。\n")
        self.status_var.set("运行中")
        self.start_button.configure(state="disabled")
        self.stop_button.configure(state="normal")
        self.reader_thread = threading.Thread(target=self._read_output, daemon=True)
        self.reader_thread.start()

    def _read_output(self) -> None:
        assert self.process is not None and self.process.stdout is not None
        for line in self.process.stdout:
            self.output_queue.put(line)
        self.output_queue.put("__PROCESS_FINISHED__\n")

    def _poll_output(self) -> None:
        try:
            while True:
                line = self.output_queue.get_nowait()
                if line == "__PROCESS_FINISHED__\n":
                    code = self.process.poll() if self.process is not None else None
                    self.status_var.set(f"已结束（退出码 {code}）")
                    self.start_button.configure(state="normal")
                    self.stop_button.configure(state="disabled")
                else:
                    self._append_log(line)
                    if line.startswith("[PROGRESS]"):
                        self.status_var.set(line.strip()[len("[PROGRESS]"):].strip())
        except queue.Empty:
            pass
        self.after(150, self._poll_output)

    def _stop(self) -> None:
        if self.process is None or self.process.poll() is not None:
            return
        if not messagebox.askyesno("确认停止", "确定停止当前任务吗？当前客户会标记为未完成。"):
            return
        self.process.terminate()
        self.status_var.set("正在停止")

    def _open_runs(self) -> None:
        runs = BASE_DIR / "runs"
        runs.mkdir(exist_ok=True)
        os.startfile(str(runs))

    def _on_close(self) -> None:
        save_input_path(Path(self.input_var.get().strip()))
        if self.process is not None and self.process.poll() is None:
            if not messagebox.askyesno("确认退出", "任务仍在运行，确定停止并退出吗？"):
                return
            self.process.terminate()
        self.destroy()


if __name__ == "__main__":
    app = RobotGui()
    app.mainloop()
