import json
import sys
import unittest
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).parents[1]))
from main import Audit, Customer, _parse_mail_timestamp, extract_invoice_email, extract_verification_code, load_config, process_customer_with_retries, resolve_mail_route


class CoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = load_config(Path(__file__).parents[1] / "config.json")

    def test_route_exact_domains(self):
        self.assertEqual(resolve_mail_route("x@erp-aid.example.invalid", self.config["mail_routes"])["key"], "aid")
        self.assertEqual(resolve_mail_route("x@example.invalid", self.config["mail_routes"])["key"], "er-helper")
        self.assertEqual(resolve_mail_route("x@erp-helper.example.invalid", self.config["mail_routes"])["key"], "er-helper")
        self.assertEqual(resolve_mail_route("x@163.com", self.config["mail_routes"])["key"], "163")
        self.assertEqual(resolve_mail_route("x@foo.ecopv0316.example.invalid", self.config["mail_routes"])["key"], "aliyun")

    def test_unknown_route_is_none(self):
        self.assertIsNone(resolve_mail_route("x@unknown.example", self.config["mail_routes"]))

    def test_invoice_email_extraction(self):
        self.assertEqual(extract_invoice_email("Invoice email\ninvoice@example.com\nEdit invoice email"), "invoice@example.com")
        self.assertEqual(extract_invoice_email("Faktura e-post\nsupport@example.com\nÄndra faktura e-post"), "support@example.com")

    def test_chinese_customer_headers(self):
        customer = Customer.from_row(2, {
            "客户中文名称": "测试客户",
            "邮箱": "test@eu-erp.example.invalid",
            "邮箱密码": "secret",
        })
        self.assertEqual(customer.name, "测试客户")
        self.assertEqual(customer.portal_email, "test@eu-erp.example.invalid")
        self.assertEqual(customer.portal_password, "secret")
        self.assertEqual(customer.mail_email, "test@eu-erp.example.invalid")
        self.assertEqual(customer.mail_password, "secret")

    def test_verification_code_extraction_uses_label(self):
        self.assertEqual(extract_verification_code("账户验证码：\n31354256\n如果没有请求获取验证码"), "31354256")
        self.assertIsNone(extract_verification_code("今天的订单号是 31354256"))

    def test_mail_timestamp_labels(self):
        reference = datetime(2026, 9, 17, 17, 20)
        self.assertEqual(_parse_mail_timestamp("NPA kundportal\n今天 17:18\n账户验证码", reference), datetime(2026, 9, 17, 17, 18))
        self.assertEqual(_parse_mail_timestamp("NPA kundportal\nToday 17:18\n账户验证码", reference), datetime(2026, 9, 17, 17, 18))
        self.assertEqual(_parse_mail_timestamp("NPA kundportal\nYesterday 23:05\n账户验证码", reference), datetime(2026, 9, 16, 23, 5))
        self.assertEqual(_parse_mail_timestamp("NPA kundportal\nTue 18:47\n账户验证码", reference), datetime(2026, 9, 15, 18, 47))
        self.assertEqual(_parse_mail_timestamp("NPA kundportal\n2026-09-16 23:05\n账户验证码", reference), datetime(2026, 9, 16, 23, 5))
        self.assertEqual(_parse_mail_timestamp("NPA kundportal\n9月17日 17:19\n账户验证码", reference), datetime(2026, 9, 17, 17, 19))
        self.assertEqual(_parse_mail_timestamp("时间：2026年9月18日 09:36 (星期五)", reference), datetime(2026, 9, 18, 9, 36))

    def test_audit_writes_explicit_results_and_failure_reasons(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is required for workbook output tests")

        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_path = root / "customers.csv"
            input_path.write_text(
                "customer_name,portal_email,portal_password,mail_email,mail_password\n"
                "Already exists,client1@example.test,,client1@example.test,secret\n"
                "Read code failed,client2@example.test,,client2@example.test,secret\n",
                encoding="utf-8",
            )
            audit = Audit(root / "run", input_path, total=2)
            first = Customer.from_row(2, {"customer_name": "Already exists", "portal_email": "client1@example.test"})
            second = Customer.from_row(3, {"customer_name": "Read code failed", "portal_email": "client2@example.test"})
            audit.result(first, status="completed", admin_action="already_exists", invoice_action="already_correct")
            audit.result(second, status="failed", failed_step="读取验证码", error="未找到登录后收到的有效验证码邮件")

            result_book = load_workbook(audit.results_xlsx_path, read_only=True, data_only=True)
            try:
                sheet = result_book.active
                headers = [cell.value for cell in sheet[1]]
                rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
            finally:
                result_book.close()
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["status"], "completed")
            self.assertEqual(rows[0]["status_label"], "已完成")
            self.assertEqual(rows[0]["admin_result"], "已存在，跳过添加")
            self.assertEqual(rows[1]["status"], "failed")
            self.assertEqual(rows[1]["failed_step"], "读取验证码")
            self.assertEqual(rows[1]["failure_reason"], "未找到登录后收到的有效验证码邮件")
            self.assertEqual(rows[1]["retryable"], "是")

            failed_book = load_workbook(audit.failed_xlsx_path, read_only=True, data_only=True)
            try:
                failed_sheet = failed_book.active
                failed_headers = [cell.value for cell in failed_sheet[1]]
                failed_rows = [dict(zip(failed_headers, row)) for row in failed_sheet.iter_rows(min_row=2, values_only=True)]
            finally:
                failed_book.close()
            self.assertEqual(len(failed_rows), 1)
            self.assertEqual(failed_rows[0]["status"], "failed")
            self.assertNotIn("Already exists", str(failed_rows[0].values()))

    def test_customer_retry_succeeds_before_failed_output(self):
        class FakeContext:
            def close(self):
                pass

        class FakeBrowser:
            def new_context(self):
                return FakeContext()

        class FakeBot:
            current_step = "读取验证码"

            def __init__(self):
                self.calls = 0

            def process(self, customer, context, route):
                self.calls += 1
                if self.calls < 3:
                    return {
                        "status": "failed",
                        "failed_step": "读取验证码",
                        "error": f"第 {self.calls} 次没有找到新邮件",
                    }
                return {
                    "status": "completed",
                    "admin_action": "already_exists",
                    "invoice_action": "already_correct",
                }

        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_path = root / "customers.csv"
            input_path.write_text("customer_name,portal_email\nRetry customer,client@example.test\n", encoding="utf-8")
            audit = Audit(root / "run", input_path, total=1)
            customer = Customer.from_row(2, {"customer_name": "Retry customer", "portal_email": "client@example.test"})
            bot = FakeBot()
            result = process_customer_with_retries(bot, customer, FakeBrowser(), {}, audit)
            self.assertEqual(bot.calls, 3)
            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["attempts"], 3)
            audit.result(customer, **result)
            self.assertFalse(audit.failed_xlsx_path.stat().st_size == 0)

    def test_customer_retry_writes_failure_only_after_three_attempts(self):
        class FakeContext:
            def close(self):
                pass

        class FakeBrowser:
            def new_context(self):
                return FakeContext()

        class AlwaysFailBot:
            current_step = "打开 Settings"

            def __init__(self):
                self.calls = 0

            def process(self, customer, context, route):
                self.calls += 1
                return {
                    "status": "failed",
                    "failed_step": "打开 Settings",
                    "error": "页面未找到 Settings",
                }

        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_path = root / "customers.csv"
            input_path.write_text("customer_name,portal_email\nFailed customer,client@example.test\n", encoding="utf-8")
            audit = Audit(root / "run", input_path, total=1)
            customer = Customer.from_row(2, {"customer_name": "Failed customer", "portal_email": "client@example.test"})
            bot = AlwaysFailBot()
            result = process_customer_with_retries(bot, customer, FakeBrowser(), {}, audit)
            self.assertEqual(bot.calls, 3)
            self.assertEqual(result["status"], "failed")
            self.assertEqual(result["attempts"], 3)
            self.assertIn("连续尝试 3 次", result["error"])
            audit.result(customer, **result)
            from openpyxl import load_workbook
            failed_book = load_workbook(audit.failed_xlsx_path, read_only=True, data_only=True)
            try:
                rows = list(failed_book.active.iter_rows(min_row=2, values_only=True))
            finally:
                failed_book.close()
            self.assertEqual(len(rows), 1)


if __name__ == "__main__":
    unittest.main()
